# -*- coding: utf-8 -*-
"""把 doc.md 转成《宿舍求生》玩法设计方案.docx（并同步输出 .md）"""
import os, re, shutil, sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(HERE, "doc.md")
OUT_DOCX = os.path.join(ROOT, "《宿舍求生》玩法设计方案.docx")
OUT_MD = os.path.join(ROOT, "《宿舍求生》玩法设计方案.md")

FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x47, 0x55, 0x69)

# ---------------- 基础工具 ----------------
def set_cn(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)

def hrule(p):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), "C9D3E0")
    bd.append(b); pPr.append(bd)

def left_bar(p, color="8FA8C4"):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:left")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18")
    b.set(qn("w:space"), "8"); b.set(qn("w:color"), color)
    bd.append(b); pPr.append(bd)

INLINE = re.compile(r"\*\*(.+?)\*\*")

def add_runs(p, text, size=10, color=None, bold_all=False):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            set_cn(p.add_run(text[pos:m.start()]), size, bold_all, color)
        set_cn(p.add_run(m.group(1)), size, True, color)
        pos = m.end()
    if pos < len(text):
        set_cn(p.add_run(text[pos:]), size, bold_all, color)
    if not text:
        set_cn(p.add_run(""), size, bold_all, color)

# ---------------- 文档设置 ----------------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.8)
USABLE = 21.0 - 3.2   # cm

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.28

lines = open(SRC, encoding="utf-8").read().split("\n")

def add_heading(text, level):
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt({1: 0, 2: 18, 3: 12, 4: 9}[level])
    pf.space_after = Pt({1: 8, 2: 7, 3: 5, 4: 4}[level])
    if level == 1:
        pf.space_before = Pt(0)
    add_runs(p, text, size=sizes[level], color=ACCENT, bold_all=True)
    if level <= 2:
        hrule(p)
    return p

def add_table(block):
    rows = []
    for ln in block:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{2,}:?", c or "-") for c in rows[1]):
        rows.pop(1)
    ncol = max(len(r) for r in rows)
    weights = []
    for j in range(ncol):
        lens = [len(r[j]) for r in rows if j < len(r)]
        weights.append(max(6, sum(lens) / max(1, len(lens)) ** 0.5 if lens else 6))
    weights = [max(1.0, w) for w in weights]
    tot = sum(weights)
    widths = [USABLE * w / tot for w in weights]

    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, r in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            cell = cells[j]
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.06
            if i == 0:
                add_runs(p, txt, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF), bold_all=True)
                shade(cell, "1F3A5F")
            else:
                add_runs(p, txt, size=8.5, color=None)
                if i % 2 == 0:
                    shade(cell, "F7FAFC")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

i = 0
while i < len(lines):
    ln = lines[i]
    s = ln.strip()

    if not s:
        i += 1; continue

    # 水平线
    if re.fullmatch(r"-{3,}", s):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); hrule(p)
        i += 1; continue

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        add_heading(m.group(2).strip(), len(m.group(1)))
        i += 1; continue

    # 表格
    if s.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i]); i += 1
        add_table(block); continue

    # 引用块
    if s.startswith(">"):
        grp = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            grp.append(lines[i].strip()[1:].strip()); i += 1
        for g in grp:
            if not g:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, g, size=10, color=GREY, bold_all=False)
            left_bar(p)
        continue

    # 无序列表
    if s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.75)
        add_runs(p, s[2:], size=10)
        i += 1; continue

    # 有序列表
    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.75)
        add_runs(p, m.group(2), size=10)
        i += 1; continue

    # 普通段落
    p = doc.add_paragraph()
    add_runs(p, s, size=10)
    i += 1

doc.save(OUT_DOCX)
shutil.copyfile(SRC, OUT_MD)
print("OK ->", OUT_DOCX)
print("OK ->", OUT_MD)
print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))
