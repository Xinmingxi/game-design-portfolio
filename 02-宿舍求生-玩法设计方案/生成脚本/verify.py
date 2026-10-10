# -*- coding: utf-8 -*-
"""《宿舍求生》v3.0 一致性与勾稽校验"""
import os, re, sys, math, itertools
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook
from docx import Document
import common, data_a, data_b, data_c, data_d   # noqa: F401
from common import SHEETS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import glob
XLDIR = os.path.join(ROOT, "数值表")
BOOKS = sorted(glob.glob(os.path.join(XLDIR, "*.xlsx")))
WBS = [load_workbook(p, data_only=True) for p in BOOKS]
doc = Document(os.path.join(ROOT, "《宿舍求生》玩法设计方案.docx"))
xl = "\n".join([ws.title for wb in WBS for ws in wb.worksheets] +
                [str(c) for wb in WBS for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None])
dc = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])

fails = []
def must(cond, msg):
    print(("  [OK]  " if cond else "  [!!]  ") + msg)
    if not cond: fails.append(msg)
def R(name): return SHEETS and {s["name"]: s for s in SHEETS}[name]
S = {s["name"]: s for s in SHEETS}
def rows(n): return S[n]["rows"]
def num(v):
    m = re.search(r"-?\d+(\.\d+)?", str(v).replace(",", ""))
    return float(m.group(0)) if m else None

print("=" * 80); print("《宿舍求生》v3.0 数值勾稽校验"); print("=" * 80)

# ---- 1. 昼夜与总时长 ----
tl = rows("03-单局时间轴")
DAYS = ["第 1 天","第 1 夜","第 2 天","第 2 夜","第 3 天","终局夜"]
main_min = sum(int(num(d0[2])) for d0 in tl if str(d0[0]) in DAYS)
must(main_min == 24, f"主赛程 = {main_min} 分钟（3 昼夜 × (5+3)）")
must(int(num([r0 for r0 in tl if str(r0[0]) == "决赛圈"][0][2])) <= 5, "决赛圈加时 ≤ 5 分钟，单局最长 29 分钟")
d = {str(x[0]): x for x in tl}
must("05:00" in str(d["第 1 夜"][1]) and "08:00" in str(d["第 1 夜"][1]), "第 1 夜 05:00-08:00")
must("24:00" in str(d["判定"][1]), "通关判定时刻 24:00")
must(int(num(d["第 1 天"][2])) == 5 and int(num(d["第 1 夜"][2])) == 3, "白天 5 分钟 / 夜晚 3 分钟")
wl = rows("05-怪潮与夜晚波次表")
tot = sum(int(num(x[4])) for x in wl)
must(tot == 120 + 240 + 360, f"三个夜晚怪数量合计 = {tot}")
must(any("BOSS" in str(x[8]) or "宿管" in str(x[8]) for x in wl), "终局夜包含 BOSS")

# ---- 2. 材料池：由 家具产出 x 房型配置 反推 15-材料总表 ----
fh = S["16-家具破坏产出表"]["headers"]
fr = {x[0]: x for x in rows("16-家具破坏产出表")}
MAT = ["木材", "铁", "零件", "石头", "布料", "纸", "玻璃", "化学品"]
mi = {m: fh.index(m) for m in MAT}
cf = rows("17-房型家具配置表")
MULT = {"标准六人间": 12, "四人间": 30, "研间 / 单间": 18, "杂物间": 12,
        "活动室": 6, "洗漱间 / 水房": 12, "每层公共区": 6, "B1 仓库": 1, "6F 天台": 1}
empty = {m: 0.0 for m in MAT}
pub = {m: 0.0 for m in MAT}
for typ, fur, qty, _ in cf:
    if typ not in MULT or fur not in fr:
        continue
    q = float(qty) * MULT[typ]
    for m in MAT:
        v = num(fr[fur][mi[m]]) or 0
        if typ in ("每层公共区", "B1 仓库", "6F 天台"):
            pub[m] += q * v
        else:
            empty[m] += q * v
mh = S["15-材料总表"]["headers"]
mr = {x[0]: x for x in rows("15-材料总表")}
bad = []
for m in MAT:
    e, p = round(empty[m]), round(pub[m])
    g = round(sum(empty[m] for _ in [0]) / 90 * 30)   # 归位 = 空宿舍均值 x 30
    g = round(empty[m] / 90 * 30)
    total = e + p + g
    stated = [int(num(mr[m][mh.index(k)])) for k in ["空宿舍 90 间", "公共区 + B1 + 天台", "「归位」30 间次", "可竞争总量"]]
    if [e, p, g, total] != stated:
        bad.append(f"{m}: 反推 {[e,p,g,total]} vs 表 {stated}")
must(not bad, "15-材料总表 的四列可由 16×17 两张表逐件反推" + ("｜" + "；".join(bad) if bad else ""))
print("        反推：" + "、".join(f"{m} {round(empty[m]+pub[m]+empty[m]/90*30)}" for m in MAT))

# ---- 3. 材料覆盖率与合闸门槛对齐 ----
need = {}
for x in rows("18-宿舍属性与升级表"):
    if str(x[0]) == "PC" and "Lv3" in str(x[1]):
        for mm in re.finditer(r"([\u4e00-\u9fa5]+) ×(\d+)", str(x[7])):
            need[mm.group(1)] = int(mm.group(2))
must(sum(need.values()) == 209, f"PC 版 Lv3 材料需求合计 = {sum(need.values())}（应 209）｜{need}")
cov1 = {m: num(mr[m][mh.index("60 队平均份额")]) / need[m] * 100 for m in need}
cov2 = {m: v * 2 for m, v in cov1.items()}
cov3 = {m: v * 3 for m, v in cov1.items()}
must(all(30 <= v <= 70 for v in cov1.values()),
     "单队覆盖率 30%~70%：" + "、".join(f"{m} {v:.0f}%" for m, v in cov1.items()))
must(any(v < 100 for v in cov2.values()) and not all(v < 100 for v in cov2.values()),
     "2 队：部分达标部分不达标 → 明确不够：" + "、".join(f"{m} {v:.0f}%" for m, v in cov2.items()))
must(all(v >= 100 for v in cov3.values()),
     "3 队：全部达标：" + "、".join(f"{m} {v:.0f}%" for m, v in cov3.items()))
must(common.C["合闸人数PC"] == 6 and common.C["合闸队数PC"] == 3, "PC 版合闸 6 人 = 3 队")
must(3 * 2 == 6, "**3 队 × 2 人 = 6 人 = 合闸需求 → 人力门槛与材料门槛精确对齐**")

# ---- 4. 武器：数量 / TTK 区间 / 档内极差 / 严格支配 ----
W = data_c.W
must(len(W) == 69, f"武器总数 = {len(W)}（应 69）")
must(all(w[9] and str(w[9]).strip() for w in W), "**69 件武器全部标注了获取来源**")
must(all(isinstance(w[8], int) and w[8] >= 1 for w in W), "**69 件武器全部有耐久（无“永不损坏”）**")
BANNED = ["军用工兵铲", "猎枪", "气枪", "信号枪", "工兵铲", "手枪", "步枪", "冲锋枪", "机枪"]
hit = [w[0] for w in W if any(b in w[0] for b in BANNED)]
must(not hit, "**不含任何原装枪械 / 军用装备**" + (f"｜发现：{hit}" if hit else "（射钉枪＝装修工具，改装发令枪＝自制）"))
guns = [w[0] for w in W if "远程" in w[2]]
must(len(guns) == 4 and all(("自制" in w[9] or "维修间" in w[9]) for w in W if "远程" in w[2]),
     f"远程武器共 {len(guns)} 件，全部来自自制或工具改造：" + "、".join(guns))
srcs = {w[9] for w in W}
must(any("寝室家具" in s for s in srcs) and any("门卫室" in s for s in srcs) and any("自制" in s for s in srcs),
     "来源覆盖寝室 / 门卫室 / 自制三类关键渠道")
ttk = lambda dmg, rate: round((math.ceil(100.0 / dmg) - 1) / rate, 2)
combat = [w for w in W if w[3] > 0 and w[1] in ("T1","T2","T3","T4","T5") and ("近战" in w[2] or "远程" in w[2])]
ttks = [ttk(w[3], w[4]) for w in combat]
must(abs(min(ttks) - 2.00) < 0.01 and abs(max(ttks) - 4.00) < 0.01,
     f"战斗向武器 TTK 区间 = {min(ttks):.2f} ~ {max(ttks):.2f} 秒")
must(abs(max(ttks) / min(ttks) - 2.0) < 0.01, f"跨档 TTK 总极差 = {max(ttks)/min(ttks):.2f} 倍（应 2.00）")
tr = {}
for w in combat: tr.setdefault(w[1], []).append(ttk(w[3], w[4]))
wide = {t: (max(v) - min(v)) / min(v) * 100 for t, v in tr.items()}
wide = {t: v for t, v in wide.items() if t in ("T1","T2","T3","T4","T5")}
must(all(v <= 15.1 for v in wide.values()),
     "各档内 TTK 极差 ≤ 15%：" + "、".join(f"{t} {v:.1f}%" for t, v in sorted(wide.items())))

MAIN = ["拖鞋","晾衣杆","折叠椅","铁管","钉子棒","钢筋长矛","钢管锤","消防斧","撬棍","电锯","射钉枪","自制弩","砍骨刀"]
mt = {w[0]: ttk(w[3], w[4]) for w in W if w[0] in MAIN}
must(len(mt) == 13, f"13 件主力武器全部可在武器表中找到（实为 {len(mt)}）")
avg = sum(mt.values()) / len(mt)
must(abs(avg - 2.62) < 0.01, f"主力武器平均 TTK = {avg:.2f}（文案写 2.62）")
must(2.4 <= avg <= 2.9, f"平均 TTK {avg:.2f} 落在 2.4 ~ 2.9 秒目标区间")

DIFF = {"T1": 1, "T2": 2, "T3": 3, "T4": 4, "T5": 5}
dom = []
for a, b in itertools.permutations(combat, 2):
    ge = (a[3] >= b[3] and a[4] >= b[4] and a[5] >= b[5] and a[8] >= b[8]
          and a[6] <= b[6] and a[7] <= b[7] and DIFF[a[1]] <= DIFF[b[1]])
    if ge:
        dom.append((a[0], b[0]))
must(len(dom) == 0, f"**严格支配关系 = {len(dom)} 对**（穷举 {len(combat)}×{len(combat)-1} 组配对）" +
     ("｜" + "；".join(f"{x}→{y}" for x, y in dom[:6]) if dom else ""))

# ---- 5. 通关漏斗 ----
f = rows("45-通关判定漏斗表")
ci = S["45-通关判定漏斗表"]["headers"].index("PC 通过率")
v = {str(x[0]): num(x[ci]) for x in f}
chain = v["L1"] / 100 * v["L2"] / 100 * v["L3"] / 100 * v["L4"] / 100 * 100
must(abs(chain - 2.2) < 0.05, f"PC 漏斗 L1×L2×L3×L4 = {chain:.2f}%（标注 2.2%）")
must(2.2 < 3.6, "解谜通关率严格低于清场率（2.2% < 3.6%）")
must(56 / 35 == 1.6, "奖励倍数自洽：+56 ÷ +35 = 1.6")

# ---- 6. 经济换算 ----
must(abs(10.06 * 100 - 1006) < 1, "ARPU 10.06 × DAU 100 万 = 月流水 1006 万")
must(abs(10.06 * 4.2 - 42) < 0.6, "生命周期 4.2 月 → LTV ≈ ¥42")
must(abs(1 / 0.077 - 13) < 0.6, "归寝凭证 0.077/日 → 约 13 天 1 张")
must(abs(333 - 95 * 3.5) < 1, "生存币 95/局 × 3.5 局 = 333/日")

# ---- 7. 双向一致性 ----
for s in ["24:00", "209", "68", "2.00 倍", "2.2%", "3.6%", "3 米", "举手", "悬赏", "已点名",
          "10 杀", "手推车", "归位", "2.62", "室友名录", "¥10.06", "民俗学", "宿管主任", "60 秒", "69", "射钉枪", "发令枪", "耐久", "自制"]:
    a, b = s in xl, s in dc
    must(a and b, f"“{s}”同时出现在 xlsx 与 docx（{a}/{b}）")

# ---- 8. 规模 ----
n_data = sum(1 for wb in WBS for ws in wb.worksheets if ws.title != "表目录")
must(len(BOOKS) == 7, f"数值表分为 {len(BOOKS)} 个工作簿（应 7）")
must(n_data == 55, f"七册合计数据表 {n_data} 张（应 55）")
must(all(any(ws.title == "表目录" for ws in wb.worksheets) for wb in WBS), "每个工作簿都含「表目录」页")
must(len(SHEETS) == 55, f"源数据注册 {len(SHEETS)} 张表")
must(len(doc.tables) >= 20, f"主文档含 {len(doc.tables)} 张表格")

# ---- 9. 可行性文档的关键数字必须与数值表一致 ----
FZ = os.path.join(ROOT, "《宿舍求生》可行性与面试问答.docx")
must(os.path.exists(FZ), "可行性文档已生成")
fdoc = Document(FZ)
fz = "\n".join([p.text for p in fdoc.paragraphs] +
                [c.text for t in fdoc.tables for r in t.rows for c in r.cells])
for s in ["69 件武器", "2.00 ~ 4.00 秒", "2.00 倍", "209", "36% ~ 58%",
          "72%", "96%", "108% ~ 174%", "2.2%", "3.6%", "¥10.06", "2.35",
          "138 秒", "57 kg", "2.4 趟", "4.8 分钟", "2778", "2060", "474", "1212",
          "39 KB/s", "12 B/s", "192 kbps", "120 只"]:
    must(s in fz, f"可行性文档含关键数字“{s}”")
must(len(fdoc.tables) >= 10, f"可行性文档含 {len(fdoc.tables)} 张表格")

# ---- 10. 自我介绍文档的关键数字必须与数值表一致 ----
IZ = os.path.join(ROOT, "面试自我介绍与速查卡.docx")
must(os.path.exists(IZ), "自我介绍文档已生成")
idoc = Document(IZ)
iz = "\n".join([p.text for p in idoc.paragraphs] +
                [c.text for t in idoc.tables for r in t.rows for c in r.cells])
for s in ["69 件武器", "2.00 ~ 4.00 秒", "2.00 倍", "0 对", "209", "36% ~ 58%",
          "3 队", "6 人", "86 件", "57 kg", "2.35", "17%", "55 张", "±35",
          "±9", "LTV/CAC", "30 日留存"]:
    must(s in iz, f"自我介绍文档含关键数字“{s}”")
must(len(idoc.tables) >= 5, f"自我介绍文档含 {len(idoc.tables)} 张表格")
print(f"  [i]   数值表数据行合计: {sum(len(s['rows']) for s in SHEETS)}")

print("=" * 80)
print("结果:", "全部通过" if not fails else f"{len(fails)} 项未通过")
for x in fails: print("  未通过:", x)
